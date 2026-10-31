import json
from io import BytesIO
from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient
from google.genai import types
from openpyxl import Workbook

from app.config import settings
from app.document_reader import xlsx_sheets_to_text
from app.exceptions import NoItemsFoundError, TooManySheetsError
from app.main import app, get_extractor
from app.models.input_schemas import InputDocument
from app.models.output_schemas import Questionnaire, SheetResult
from app.questionnaire_extractor import QuestionnaireExtractor
from tests.conftest import DATA_DIR, EXAMPLES_DIR

SPEC_EXAMPLE = (DATA_DIR / "spec_example.json").read_text(encoding="utf-8")
XLSX = (EXAMPLES_DIR / "02-tipo-xlsx" / "cuestionario.xlsx").read_bytes()


def make_workbook(sheets: dict[str, list[list]]) -> bytes:
    """Crea un .xlsx en memoria: {nombre de hoja: filas}."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(name)
        for row in rows:
            sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


class FakeModels:
    """Simula client.aio.models: devuelve las respuestas en el orden de las llamadas."""

    def __init__(self, texts: list[str | None]):
        self.responses = [
            SimpleNamespace(text=text, candidates=[SimpleNamespace(finish_reason=types.FinishReason.STOP)])
            for text in texts
        ]
        self.calls = []

    async def generate_content(self, **kwargs):
        self.calls.append(kwargs)
        return self.responses[len(self.calls) - 1]


def make_extractor(texts: list[str | None]):
    models = FakeModels(texts)
    client = SimpleNamespace(aio=SimpleNamespace(models=models))
    return QuestionnaireExtractor(client=client), models


def xlsx_document(content: bytes) -> InputDocument:
    return InputDocument(filename="cuestionarios.xlsx", format="xlsx", content=content)


def test_reads_every_sheet_of_example():
    sheets = xlsx_sheets_to_text(XLSX)

    assert [name for name, _ in sheets] == ["Cuestionario_ES", "Cuestionario_IT", "Cuestionario_EN"]
    assert sheets[1][1].startswith("Questionario di esempio")


async def test_sends_each_sheet_in_its_own_call():
    extractor, models = make_extractor([SPEC_EXAMPLE] * 3)

    results = await extractor.extract_sheets(xlsx_document(XLSX))

    assert len(models.calls) == 3
    assert [result.sheet_name for result in results] == ["Cuestionario_ES", "Cuestionario_IT", "Cuestionario_EN"]
    assert all(result.questionnaire == Questionnaire.model_validate_json(SPEC_EXAMPLE) for result in results)
    assert all(result.error is None for result in results)


async def test_error_in_one_sheet_does_not_affect_others():
    extractor, _ = make_extractor([SPEC_EXAMPLE, "no es json", SPEC_EXAMPLE])

    results = await extractor.extract_sheets(xlsx_document(XLSX))

    assert results[0].questionnaire is not None
    assert results[1].questionnaire is None
    assert results[1].error.error == "respuesta_invalida"
    assert results[2].questionnaire is not None


async def test_skips_empty_sheets():
    content = make_workbook({"Vacía": [], "Cuestionario": [["1", "¿Pregunta?"]]})
    extractor, models = make_extractor([SPEC_EXAMPLE])

    results = await extractor.extract_sheets(xlsx_document(content))

    assert len(models.calls) == 1
    assert [result.sheet_name for result in results] == ["Cuestionario"]


async def test_raises_when_all_sheets_are_empty():
    extractor, models = make_extractor([])

    with pytest.raises(NoItemsFoundError):
        await extractor.extract_sheets(xlsx_document(make_workbook({"A": [], "B": []})))
    assert models.calls == []


async def test_raises_when_too_many_sheets():
    sheets = {f"Hoja{i}": [["contenido"]] for i in range(settings.MAX_SHEETS + 1)}
    extractor, models = make_extractor([])

    with pytest.raises(TooManySheetsError):
        await extractor.extract_sheets(xlsx_document(make_workbook(sheets)))
    assert models.calls == []


class FakeSheetsExtractor:
    async def extract_sheets(self, document):
        return [SheetResult(sheet_name="Cuestionario_ES", questionnaire=Questionnaire.model_validate_json(SPEC_EXAMPLE))]


@pytest.fixture
def client():
    app.dependency_overrides[get_extractor] = FakeSheetsExtractor
    yield TestClient(app)
    app.dependency_overrides.clear()


def test_endpoint_returns_list_with_spanish_field_names(client):
    response = client.post("/v1/questionnaires/extract-sheets", files={"file": ("cuestionarios.xlsx", XLSX)})

    assert response.status_code == 200
    [result] = response.json()
    assert set(result) == {"hoja", "cuestionario", "error"}
    assert result["hoja"] == "Cuestionario_ES"
    assert result["cuestionario"]["items"][0]["enunciado"] == json.loads(SPEC_EXAMPLE)["items"][0]["enunciado"]


def test_endpoint_rejects_non_xlsx(client):
    docx = (EXAMPLES_DIR / "01-tipo-doc" / "cuestionario.docx").read_bytes()

    response = client.post("/v1/questionnaires/extract-sheets", files={"file": ("cuestionario.docx", docx)})

    assert response.status_code == 415
    assert response.json()["error"] == "formato_invalido"
